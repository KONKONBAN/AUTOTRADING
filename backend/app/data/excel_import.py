"""Nhập dữ liệu từ file Excel/CSV do người dùng tải lên.

1. BCTC (Báo cáo tài chính) - hỗ trợ:
   - File mẫu của ứng dụng (có cột "key"), và
   - File xuất từ Vietstock / CafeF / FiinPro / VNDirect / SSI... dạng "chỉ tiêu × năm":
     tự nhận diện dòng chỉ tiêu theo tên tiếng Việt/tiếng Anh, cột năm và đơn vị
     (đồng / nghìn đồng / triệu đồng / tỷ đồng) rồi quy đổi về tỷ đồng.
2. Lịch sử giá - định dạng Amibroker/MetaStock (<Ticker>,<DTYYYYMMDD>,...) hoặc bảng có tiêu đề
   (Ngày, Giá mở cửa, Giá cao nhất, Giá thấp nhất, Giá đóng cửa, Khối lượng...).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..db import get_db
from . import fundamentals


def norm(text) -> str:
    """Chuẩn hóa nhãn: bỏ dấu tiếng Việt, chữ thường, bỏ đánh số đầu dòng."""
    s = unicodedata.normalize("NFD", str(text or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()
    s = re.sub(r"\(.*?\)", " ", s)
    s = re.sub(r"^[\s\-–•*+]*((\d+(\.\d+)*|[ivxlc]+|[a-h])[\.\)/:]\s*)+", "", s.strip())
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9%/ ]+", " ", s)).strip()


_VN_NUM = re.compile(r"^[-(]?\d{1,3}(\.\d{3})+(,\d+)?\)?$|^[-(]?\d+,\d{1,2}\)?$")
_EN_NUM = re.compile(r"^[-(]?\d{1,3}(,\d{3})+(\.\d+)?\)?$|^[-(]?\d+\.\d{1,2}\)?$")
# Số thập phân KHÔNG thể là nhóm hàng nghìn: phần nguyên ≥ 4 chữ số hoặc bằng 0, hoặc ≥ 4 chữ số thập phân
_EN_DEC = re.compile(r"^[-(]?(\d{4,}|0)\.\d+\)?$|^[-(]?\d+\.\d{4,}\)?$")
_VN_DEC = re.compile(r"^[-(]?(\d{4,}|0),\d+\)?$|^[-(]?\d+,\d{4,}\)?$")
# Nhóm hàng nghìn hợp lệ: 1-3 chữ số đầu (không bắt đầu bằng 0) rồi các nhóm 3 chữ số
_GROUP_DOT = re.compile(r"^[1-9]\d{0,2}(\.\d{3})+$")
_GROUP_COMMA = re.compile(r"^[1-9]\d{0,2}(,\d{3})+$")


def detect_convention(grid: list[list], delimiter: str | None = None) -> str | None:
    """Đoán quy ước số của cả file: 'vn' (1.234.567,8) hay 'en' (1,234,567.8)."""
    vn = en = 0
    for row in grid[:400]:
        for c in row:
            if not isinstance(c, str):
                continue
            t = c.strip().replace("\xa0", "").replace(" ", "")
            if _VN_NUM.match(t) or _VN_DEC.match(t):
                vn += 1
            elif _EN_NUM.match(t) or _EN_DEC.match(t):
                en += 1
    if delimiter == ";":
        vn += 2
    if vn > en:
        return "vn"
    if en > vn:
        return "en"
    return None


def _num(v, conv: str | None = None) -> float | None:
    """Đọc số từ ô Excel/CSV: hỗ trợ 1.234.567 | 1,234,567 | 1.234,5 | 1,234.5 | (123) = -123.

    conv = 'vn' (dấu chấm phân cách hàng nghìn, phẩy thập phân) | 'en' | None (tự đoán từng ô).
    """
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    if not s or s in {"-", "--", "n/a", "N/A"}:
        return None
    neg = s.startswith("(") and s.endswith(")") or s.startswith("-")
    s = s.strip("()").lstrip("-")
    if "," in s and "." in s:
        # dấu xuất hiện sau cùng là dấu thập phân: 1.234,5 (VN) | 1,234.5 (EN)
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        if s.count(",") > 1 or (conv != "vn" and _GROUP_COMMA.match(s)):
            s = s.replace(",", "")      # 1,234,567 | 1,234 (EN / không rõ)
        else:
            s = s.replace(",", ".")     # 1,5 | 0,125 | 1,234 (VN = một phẩy hai ba bốn)
    elif "." in s:
        if s.count(".") > 1 or (conv != "en" and _GROUP_DOT.match(s)):
            s = s.replace(".", "")      # 1.234.567 | 12.345 (VN / không rõ: mười hai nghìn)
        # còn lại là số thập phân: 1.5 | 0.125 | 44010.123 (phần nguyên ≥ 4 chữ số không thể là nhóm nghìn)
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


# --------------------------------------------------------------------------- BCTC
# (trường, [mẫu ưu tiên], [loại trừ], cộng dồn?)
FIELD_PATTERNS: list[tuple[str, list[str], list[str], bool]] = [
    ("revenue", ["doanh thu thuan ve ban hang va cung cap dich vu", "doanh thu thuan", "net revenue",
                 "net sales", "tong thu nhap hoat dong", "doanh thu hoat dong", "doanh thu"],
     ["tai chinh", "giam tru", "chua thuc hien", "khac", "tang truong", "%"], False),
    ("gross_profit", ["loi nhuan gop ve ban hang", "loi nhuan gop", "gross profit"], ["%", "bien"], False),
    ("interest_expense", ["trong do chi phi lai vay", "chi phi lai vay", "interest expense",
                          "chi phi lai va cac chi phi tuong tu"], ["%", "tra"], False),
    ("pretax_income", ["tong loi nhuan ke toan truoc thue", "loi nhuan truoc thue", "profit before tax",
                       "pre tax profit"], ["%"], False),
    ("net_income", ["loi nhuan sau thue cua co dong cua cong ty me", "loi nhuan sau thue cong ty me",
                    "loi nhuan sau thue cua co dong cua ngan hang me", "co dong cua ngan hang me",
                    "co dong cua cong ty me", "loi nhuan sau thue cua co dong cong ty me",
                    "net profit attributable to", "loi nhuan sau thue thu nhap doanh nghiep",
                    "loi nhuan sau thue", "net profit after tax", "profit after tax", "net income"],
     ["chua phan phoi", "%", "khong kiem soat", "thieu so", "tang truong", "phan phoi"], False),
    ("depreciation", ["khau hao tai san co dinh", "khau hao tscd", "khau hao", "depreciation"], ["luy ke", "hao mon"], False),
    ("total_assets", ["tong cong tai san", "tong tai san", "total assets"], ["%", "ngan han", "dai han"], False),
    ("current_assets", ["tai san ngan han", "current assets", "tai san luu dong"], ["khac", "%", "dau tu"], False),
    ("cash", ["tien va cac khoan tuong duong tien", "cash and cash equivalents", "tien va tuong duong tien"],
     ["luu chuyen", "cuoi ky", "dau ky", "anh huong"], False),
    ("inventory", ["hang ton kho", "inventories", "inventory"], ["du phong", "giam gia", "bien dong", "tang giam"], False),
    ("total_liabilities", ["tong no phai tra", "no phai tra", "total liabilities", "tong cong no"],
     ["ngan han", "dai han", "khac", "%", "nguoi ban"], False),
    ("current_liabilities", ["no ngan han", "current liabilities"], ["khac", "%"], False),
    ("total_debt", ["vay va no thue tai chinh ngan han", "vay va no thue tai chinh dai han", "vay ngan han",
                    "vay dai han", "short term borrowings", "long term borrowings", "short term debt",
                    "long term debt"], ["chi phi", "lai", "tra no", "tien thu", "tien chi"], True),
    ("equity", ["tong von chu so huu", "von chu so huu", "von va cac quy", "total equity", "owners equity",
                "shareholders equity"],
     ["khac", "%", "tren", "tong cong nguon", "ty suat", "loi nhuan tren"], False),
    ("retained_earnings", ["loi nhuan sau thue chua phan phoi", "retained earnings", "undistributed"], [], False),
    ("operating_cf", ["luu chuyen tien thuan tu hoat dong kinh doanh", "luu chuyen tien thuan tu hoat dong san xuat kinh doanh",
                      "net cash from operating", "net cash flows from operating", "luu chuyen tien te rong tu cac hoat dong san xuat kinh doanh"],
     [], False),
    ("capex", ["tien chi de mua sam xay dung tscd", "tien chi de mua sam xay dung tai san co dinh",
               "tien chi mua sam tai san co dinh", "mua sam tai san co dinh", "purchase of fixed assets",
               "purchases of property", "capital expenditure"], ["thanh ly", "nhuong ban"], False),
    ("shares", ["so luong co phieu dang luu hanh", "so luong co phieu luu hanh", "so co phieu luu hanh",
                "shares outstanding", "co phieu pho thong dang luu hanh"], ["quy", "ky tru"], False),
    ("dps", ["co tuc moi co phieu", "co tuc tien mat moi co phieu", "dividend per share", "dps"], ["%"], False),
]
OWNER_CAPITAL = ["von gop cua chu so huu", "von dau tu cua chu so huu", "von co phan", "von dieu le", "share capital"]

UNIT_PATTERNS = [("ty dong", 1.0), ("ty vnd", 1.0), ("billion", 1.0), ("trieu dong", 1e-3), ("trieu vnd", 1e-3),
                 ("million", 1e-3), ("nghin dong", 1e-6), ("ngan dong", 1e-6), ("thousand", 1e-6),
                 ("vnd", 1e-9), ("dong", 1e-9)]


# Cột quý / bán niên / lũy kế 6-9 tháng: Q1/2024, 2024Q1, 4Q2024, Quý 4/2024, H1, 2024H1, 1H2025, 6T, 9 tháng, 6M
_QUARTER = re.compile(r"(?<![a-zà-ỹ])(q|quy|quý)\s*[1-4](?!\d)|(?<!\d)[1-4]\s*q(?![a-zà-ỹ])"
                      r"|(?<![a-zà-ỹ])h[12](?!\d)|(?<!\d)[12]\s*h(?![a-zà-ỹ])|b[aá]n\s*ni[eê]n"
                      r"|(?<!\d)[69]\s*(tháng|thang|th|t|m)(?![a-zà-ỹ])", re.I)


def _year_of(v) -> int | None:
    if isinstance(v, (int, float)) and 1990 <= v <= 2100 and float(v).is_integer():
        return int(v)
    if isinstance(v, (datetime, date)):
        return v.year
    s = str(v or "")
    if _QUARTER.search(s):
        return None  # cột quý / 6 tháng / 9 tháng không phải số liệu cả năm
    if re.search(r"(19|20)\d{2}\s*[FE]\b", s, re.I) or re.search(r"du phong|ke hoach|forecast|\bKH\b", norm(s) + " " + s):
        return None  # cột dự phóng / kế hoạch
    m = re.search(r"(19|20)\d{2}", s)
    if m and len(s) <= 30:
        return int(m.group(0))
    return None


class FileFormatError(ValueError):
    pass


def _decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _grids(data: bytes, filename: str) -> list[tuple[str, list[list], str | None]]:
    """Đọc file thành các bảng [(tên sheet, dòng, quy ước số)]."""
    name = filename.lower()
    if name.endswith(".xls"):
        raise FileFormatError("File .xls (Excel 97-2003) chưa được hỗ trợ - hãy mở bằng Excel và lưu lại dạng .xlsx.")
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        raise FileFormatError("Đây là file PDF - hãy dùng mục 'AI đọc ảnh chụp / PDF báo cáo tài chính'.")
    if name.endswith((".xlsx", ".xlsm")) and data[:2] != b"PK":
        raise FileFormatError("File không phải Excel .xlsx hợp lệ - hãy mở bằng Excel và lưu lại dạng .xlsx.")
    if name.endswith((".csv", ".txt")) or data[:2] != b"PK":
        text = _decode_text(data)
        if not text.strip():
            raise FileFormatError("File rỗng.")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delim = dialect.delimiter
        except csv.Error:
            delim = max(",;\t", key=sample.count)
        try:
            rows = [row for row in csv.reader(io.StringIO(text), delimiter=delim)]
        except csv.Error as exc:
            raise FileFormatError(f"Không đọc được file CSV: {exc}") from exc
        rows = [r for r in rows if any(str(c).strip() for c in r)]
        return [("csv", rows, detect_convention(rows, delim))]
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - zip/XML/zlib hỏng… đều là lỗi định dạng file, không phải lỗi máy chủ
        raise FileFormatError("Không đọc được file Excel - hãy lưu lại dạng .xlsx hoặc CSV UTF-8.") from exc
    out = []
    try:  # read_only: lỗi XML của sheet hỏng chỉ xuất hiện khi duyệt dòng
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c not in (None, "") for c in r)]
            if rows:
                out.append((ws.title, rows, detect_convention(rows)))
    except Exception as exc:  # noqa: BLE001 - file hỏng/thiếu dữ liệu → báo lỗi rõ ràng thay vì 500
        raise FileFormatError("File Excel bị hỏng hoặc không đầy đủ - hãy mở bằng Excel và lưu lại dạng .xlsx.") from exc
    finally:
        wb.close()
    return out


def _plain(text: str) -> str:
    s = unicodedata.normalize("NFD", text).replace("đ", "d").replace("Đ", "D")
    return "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()


_MULT_BY_ZEROS = {1: (1e-6, "nghin dong"), 2: (1e-3, "trieu dong"), 3: (1.0, "ty dong")}


def _detect_unit(grid: list[list]) -> tuple[float | None, str | None]:
    """Tìm dòng 'Đơn vị: ...' (kể cả dạng '1.000 đồng', '1.000.000 VND', 'Đơn vị tính (Triệu VND)')."""
    for row in grid[:40]:
        for c in row:
            if not isinstance(c, str):
                continue
            n = _plain(c)
            if not ("don vi" in n or "unit" in n or "dvt" in n):
                continue
            m = re.search(r"1(?:[.,\s]?000)+", c)
            if m:
                zeros = m.group(0).count("000")
                if zeros in _MULT_BY_ZEROS:
                    return _MULT_BY_ZEROS[zeros]
            for pat, mult in UNIT_PATTERNS:
                if pat in n:
                    return mult, pat
    return None, None


def _row_years(row: list) -> dict[int, int]:
    """Các cột năm của một dòng tiêu đề. Ô số chỉ được coi là năm khi MỌI ô số trong dòng đều giống năm
    và các năm liền kề nhau (tránh nhầm dòng số liệu như 1.900, 2.000 thành năm)."""
    nums = [c for c in row if isinstance(c, (int, float)) and not isinstance(c, bool)]
    numeric_years = bool(nums) and all(1990 <= c <= 2100 and float(c).is_integer() for c in nums)
    cols: dict[int, int] = {}
    for j, c in enumerate(row):
        if isinstance(c, (int, float)) and not isinstance(c, bool):
            if numeric_years:
                cols[j] = int(c)
        elif (y := _year_of(c)) is not None:
            cols[j] = y
    # Tách thành các dãy năm liền nhau. Ô năm đứng lẻ (VD ô tóm tắt "CAGR 2020-2024") bị bỏ; giữ MỌI dãy
    # có ≥ 2 năm (VD 2014-2016 và 2023-2024 đều giữ, không đánh rơi các năm gần nhất);
    # toàn năm rời rạc → không phải dòng tiêu đề.
    runs: list[list[tuple[int, int]]] = []
    for j in sorted(cols):
        if runs and abs(cols[j] - runs[-1][-1][1]) <= 3:
            runs[-1].append((j, cols[j]))
        else:
            runs.append([(j, cols[j])])
    if len(runs) <= 1:
        return cols
    return {j: y for run in runs if len(run) >= 2 for j, y in run}


def _parse_grid(grid: list[list], conv: str | None = None) -> dict | None:
    # 1) tìm dòng tiêu đề chứa ≥ 2 năm
    header_idx, year_cols = None, {}
    for i, row in enumerate(grid[:60]):
        cols = _row_years(row)
        if len(cols) >= 2 or (len(cols) == 1 and len(grid) > 5):
            if len(cols) > len(year_cols):
                header_idx, year_cols = i, cols
            if len(cols) >= 2:
                break
    if not year_cols:
        return None
    # nếu một năm xuất hiện nhiều lần (đầu kỳ/cuối kỳ), giữ cột đầu tiên
    seen, cols = set(), {}
    for j, y in sorted(year_cols.items()):
        if y not in seen:
            cols[j] = y
            seen.add(y)
    first_year_col = min(cols)
    # 2) cột "key" của file mẫu
    keys = {k for k, _, _ in fundamentals.MANUAL_FIELDS}
    key_col = None
    for j in range(first_year_col):
        hits = sum(1 for r in grid[header_idx + 1:] if j < len(r) and str(r[j] or "").strip() in keys)
        if hits >= 3:
            key_col = j
            break
    lines = []
    for r in grid[header_idx + 1:]:
        label = next((str(c) for c in r[:first_year_col] if isinstance(c, str) and c.strip()), "")
        vals = {y: _num(r[j], conv) if j < len(r) else None for j, y in cols.items()}
        if not label or all(v is None for v in vals.values()):
            continue
        lines.append({"label": label, "norm": norm(label),
                      "key": str(r[key_col]).strip() if key_col is not None and key_col < len(r) else None,
                      "values": vals})
    unit_mult, unit_name = _detect_unit(grid)
    mapped: dict[str, dict[int, float]] = {}
    report = []
    if key_col is not None:
        for ln in lines:
            if ln["key"] in keys:
                mapped[ln["key"]] = ln["values"]
                report.append({"field": ln["key"], "label": ln["label"]})
        unit_mult = unit_mult or 1.0
        unit_name = unit_name or "ty dong (file mẫu)"
    else:
        used = set()
        for field, patterns, excludes, summed in FIELD_PATTERNS:
            matched = []
            for pat in patterns:
                for idx, ln in enumerate(lines):
                    if idx in used or idx in matched or any(ex in ln["norm"] for ex in excludes):
                        continue
                    if ln["norm"].startswith(pat) or (len(pat) > 12 and pat in ln["norm"]):
                        matched.append(idx)
                        if not summed:
                            break
                if matched and not summed:
                    break
            if not matched:
                continue
            if summed:
                vals = {y: sum((lines[i]["values"][y] or 0) for i in matched) for y in cols.values()}
            else:
                vals = lines[matched[0]]["values"]
            used.update(matched)
            mapped[field] = vals
            report.append({"field": field, "label": " + ".join(lines[i]["label"] for i in matched)})
        cap = next((ln for ln in lines if any(ln["norm"].startswith(p) for p in OWNER_CAPITAL)), None)
        if cap and "shares" not in mapped:
            mapped["_owner_capital"] = cap["values"]
    if len([f for f in fundamentals.REQUIRED_FIELDS if f in mapped]) < 2:
        return None
    unit_guessed = unit_mult is None
    if unit_guessed:
        # Không có dòng "Đơn vị": đoán theo độ lớn của Tổng tài sản / Doanh thu (cần người dùng xác nhận)
        ref = next((abs(v) for f in ("total_assets", "revenue", "equity") for v in (mapped.get(f) or {}).values() if v),
                   None)
        unit_mult, unit_name = ((1e-9, "dong (đoán)") if ref and ref > 1e10 else
                                (1e-6, "nghin dong (đoán)") if ref and ref > 1e7 else
                                (1e-3, "trieu dong (đoán)") if ref and ref > 1e5 else (1.0, "ty dong (đoán)"))
    years = []
    for y in sorted(set(cols.values())):
        row = {"year": y}
        for f, vals in mapped.items():
            v = vals.get(y)
            if v is None:
                continue
            if f == "shares":
                row[f] = v / 1e6 if abs(v) > 1e5 else v  # số CP → triệu CP
            elif f == "dps":
                row[f] = abs(v)
            elif f == "_owner_capital":
                row["shares"] = abs(v) * unit_mult * 0.1  # vốn góp (tỷ đồng) / 10.000đ mệnh giá → triệu CP
            elif f in {"capex", "interest_expense"}:
                row[f] = abs(v) * unit_mult
            else:
                row[f] = v * unit_mult
        if any(row.get(f) is not None for f in fundamentals.REQUIRED_FIELDS):
            years.append(row)
    missing = [f for f, _, _ in fundamentals.MANUAL_FIELDS if not any(f in y for y in years)]
    return {"years": years, "unit_detected": unit_name, "unit_guessed": unit_guessed,
            "shares_from_capital": "_owner_capital" in mapped,  # số CP suy từ vốn góp → phụ thuộc đơn vị
            "mapping": report, "missing_fields": missing}


def parse_financial_file(data: bytes, filename: str) -> dict:
    """Đọc file BCTC Excel/CSV → số liệu theo năm (tỷ đồng) để người dùng xem lại trước khi lưu."""
    best = None
    quarterly = False
    for sheet, grid, conv in _grids(data, filename):
        if any(isinstance(c, str) and _QUARTER.search(c) for row in grid[:60] for c in row):
            quarterly = True
        transposed = [list(x) for x in zip(*grid)] if grid else []
        for orientation, g in (("rows", grid), ("columns", transposed)):
            res = _parse_grid(g, conv)
            if res and (best is None or len(res["mapping"]) > len(best["mapping"])):
                best = {**res, "sheet": sheet, "orientation": orientation, "number_format": conv}
    if best is None:
        if quarterly:
            raise ValueError("File chỉ có số liệu theo QUÝ / 6 tháng. Ứng dụng phân tích BCTC NĂM - hãy dùng báo cáo "
                             "năm (hoặc cộng 4 quý cho các chỉ tiêu KQKD/lưu chuyển tiền, lấy số cuối Q4 cho bảng cân đối).")
        raise ValueError("Không nhận diện được cấu trúc BCTC trong file (cần cột năm và các dòng chỉ tiêu). "
                         "Hãy dùng file mẫu hoặc thử tính năng AI đọc file.")
    best["fields"] = [{"key": k, "label": label, "unit": unit} for k, label, unit in fundamentals.MANUAL_FIELDS]
    best["warnings"] = []
    if best.get("unit_guessed"):
        best["warnings"].append(f"Không tìm thấy dòng 'Đơn vị' trong file - đơn vị được ĐOÁN là "
                                f"{best['unit_detected'].replace(' (đoán)', '')}. Hãy kiểm tra độ lớn số liệu và dùng "
                                f"nút ×1000 / ÷1000 nếu cần.")
    if quarterly:
        best["warnings"].append("File có cả cột theo quý/6 tháng - các cột này đã được bỏ qua, chỉ lấy cột năm.")
    for y in best["years"]:
        if all(k in y for k in ("total_assets", "equity", "total_liabilities")):
            gap = abs(y["total_assets"] - y["equity"] - y["total_liabilities"])
            if gap > 0.02 * abs(y["total_assets"]):
                best["warnings"].append(f"Năm {y['year']}: Tổng tài sản ≠ Nợ phải trả + Vốn chủ sở hữu.")
    return best


def grid_as_text(data: bytes, filename: str, max_rows: int = 250, max_cols: int = 16) -> str:
    """Chuyển nội dung file thành văn bản (CSV) để AI đọc khi định dạng quá lạ."""
    parts = []
    for sheet, grid, _ in _grids(data, filename):
        buf = io.StringIO()
        w = csv.writer(buf)
        for r in grid[:max_rows]:
            w.writerow(["" if c is None else c for c in r[:max_cols]])
        parts.append(f"### Sheet: {sheet}\n{buf.getvalue()}")
    return "\n\n".join(parts)[:150_000]


# --------------------------------------------------------------------------- file mẫu
def financial_template(ticker: str | None = None) -> bytes:
    """File Excel mẫu nhập BCTC (điền sẵn số liệu hiện có nếu chọn mã)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BCTC"
    rows_existing = []
    if ticker:
        try:
            fin = fundamentals.get_financials(ticker)
            src = fin["source"]
            if src == "demo":
                rows_existing = []  # không bao giờ điền sẵn số liệu DEMO bịa ra
            elif "demo" in src:
                rows_existing = [r for r in fin["rows"] if r["year"] in fin.get("manual_years", [])]
            else:
                rows_existing = fin["rows"]
            rows_existing = rows_existing[-5:]
        except Exception:  # noqa: BLE001
            rows_existing = []
    years = [r["year"] for r in rows_existing] or list(range(date.today().year - 5, date.today().year))
    title_font = Font(bold=True, size=14, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F766E")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="CBD5E1")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3 + len(years))
    ws.cell(1, 1, f"BÁO CÁO TÀI CHÍNH NĂM - {ticker.upper() if ticker else 'MÃ CỔ PHIẾU'}").font = title_font
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="134E4A")
    ws.cell(2, 1, "Đơn vị: tỷ đồng (riêng số cổ phiếu: triệu CP; cổ tức: đồng/CP). "
                  "Không sửa cột 'key'. Có thể thêm/bớt cột năm.").font = Font(italic=True, color="475569")
    headers = ["Chỉ tiêu", "key", "Đơn vị", *years]
    for j, h in enumerate(headers, 1):
        c = ws.cell(4, j, h)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal="center")
    by_year = {r["year"]: r for r in rows_existing}
    for i, (key, label, unit) in enumerate(fundamentals.MANUAL_FIELDS, start=5):
        ws.cell(i, 1, label)
        ws.cell(i, 2, key).font = Font(color="94A3B8")
        ws.cell(i, 3, unit).font = Font(color="64748B")
        for j, y in enumerate(years, start=4):
            row = by_year.get(y, {})
            # trường được ƯỚC TÍNH hoặc bỏ trống (tính là 0) giữ nguyên ô trống - không biến thành số liệu thật
            skip = set(row.get("estimated_fields") or []) | set(row.get("defaulted_fields") or [])
            v = None if key in skip else row.get(key)
            c = ws.cell(i, j, round(v, 2) if isinstance(v, (int, float)) else None)
            c.number_format = "#,##0.00"
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if key in fundamentals.REQUIRED_FIELDS:
            ws.cell(i, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12
    for j in range(4, 4 + len(years)):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.freeze_panes = "D5"
    guide = wb.create_sheet("Huong_dan")
    for i, line in enumerate([
        "HƯỚNG DẪN NHẬP BÁO CÁO TÀI CHÍNH",
        "1. Nhập số liệu BCTC năm (ưu tiên BCTC hợp nhất đã kiểm toán) vào sheet 'BCTC', đơn vị tỷ đồng.",
        "2. Các dòng in đậm là bắt buộc: Doanh thu thuần, LNST, Tổng tài sản, Vốn chủ sở hữu.",
        "3. Chi phí lãi vay và Capex nhập số dương.",
        "4. Kiểm tra: Tổng tài sản = Nợ phải trả + Vốn chủ sở hữu.",
        "5. Lưu file và tải lên tại mục 'Dữ liệu & Nhập liệu' của ứng dụng.",
        "Ứng dụng cũng đọc được file xuất trực tiếp từ Vietstock, CafeF, FiinPro, VNDirect, SSI... (chỉ tiêu × năm).",
    ], start=1):
        guide.cell(i, 1, line).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    guide.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- giá
PRICE_HEADERS = {
    "ticker": ["ticker", "symbol", "ma ck", "ma", "ma chung khoan", "code"],
    "date": ["dtyyyymmdd", "date", "ngay", "time", "thoi gian", "trading date"],
    "open": ["open", "gia mo cua", "mo cua"],
    "high": ["high", "gia cao nhat", "cao nhat"],
    "low": ["low", "gia thap nhat", "thap nhat"],
    "adj_close": ["gia dieu chinh", "adj close", "adjusted close", "gia dc"],
    "close": ["close", "gia dong cua", "dong cua", "price", "lan cuoi", "gia"],
    "volume": ["volume", "vol", "khoi luong khop lenh", "khoi luong", "kl khop lenh", "tong khoi luong", "kl"],
}
_SUFFIX = {"k": 1e3, "m": 1e6, "b": 1e9}


def _parse_date(v, month_first: bool = False) -> str | None:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    if re.fullmatch(r"\d{8}", s):
        return datetime.strptime(s, "%Y%m%d").strftime("%Y-%m-%d")
    fmts = ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y", "%d/%m/%Y", "%d-%m-%Y") if month_first else \
        ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d", "%d.%m.%Y")
    for fmt in (*fmts, "%Y/%m/%d", "%d.%m.%Y", "%b %d, %Y", "%d %b %Y"):
        try:
            return datetime.strptime(s[:12].strip(), fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _month_first(values: list, header_row: list) -> bool:
    """Ngày dạng a/b/yyyy: tháng trước (investing.com bản tiếng Anh) hay ngày trước (Việt Nam)?
    Dựa vào bằng chứng trong file (số > 12), nếu không phân biệt được thì theo kiểu tiêu đề."""
    first_gt12 = second_gt12 = False
    for v in values:
        m = re.match(r"^\s*(\d{1,2})[/-](\d{1,2})[/-]\d{4}", str(v or ""))
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            first_gt12 |= a > 12
            second_gt12 |= b > 12
    if first_gt12 != second_gt12:
        return second_gt12
    heads = {norm(str(c)) for c in header_row if c is not None}
    return {"date", "price"} <= heads or "change %" in {str(c).strip().lower() for c in header_row if c}


def _vol(v, conv) -> float:
    """Khối lượng: 1.234.567 | 12,34M | 850.5K (investing.com)."""
    if isinstance(v, str):
        t = v.strip().replace(" ", "")
        if t and t[-1].lower() in _SUFFIX:
            base = _num(t[:-1], conv)
            return (base or 0) * _SUFFIX[t[-1].lower()]
    return _num(v, conv) or 0


def parse_price_file(data: bytes, filename: str, default_ticker: str | None = None) -> list[tuple]:
    rows_out: list[tuple] = []
    for _, grid, conv in _grids(data, filename):
        header_idx, colmap = None, {}
        for i, row in enumerate(grid[:15]):
            cm = {}
            for j, c in enumerate(row):
                n = norm(str(c).strip("<>")) if c is not None else ""
                for field, names in PRICE_HEADERS.items():
                    if field not in cm and any(n == nm or (len(nm) > 4 and n.startswith(nm)) for nm in names):
                        cm[field] = j
                        break
            if "date" in cm and ("close" in cm or "adj_close" in cm):
                header_idx, colmap = i, cm
                break
        if header_idx is None:
            # không có tiêu đề: thử định dạng Amibroker / MetaStock (cophieu68)
            colmap = {"ticker": 0, "date": 1, "open": 2, "high": 3, "low": 4, "close": 5, "volume": 6}
            header_idx = -1
        body = grid[header_idx + 1:]
        month_first = _month_first([r[colmap["date"]] for r in body if colmap["date"] < len(r)],
                                   grid[header_idx] if header_idx >= 0 else [])
        parsed: list[list] = []
        for r in body:
            try:
                tk = (str(r[colmap["ticker"]]).strip().upper() if "ticker" in colmap and colmap["ticker"] < len(r)
                      else (default_ticker or "").upper())
                d = _parse_date(r[colmap["date"]], month_first)
                raw = _num(r[colmap["close"]], conv) if "close" in colmap else None
                adj = _num(r[colmap["adj_close"]], conv) if "adj_close" in colmap else None
                c = adj if adj is not None else raw
                # file có cả giá điều chỉnh & giá gốc (CafeF): quy O/H/L về cùng hệ điều chỉnh
                f = adj / raw if (adj is not None and raw) else 1.0
                o = _num(r[colmap["open"]], conv) if "open" in colmap else None
                h = _num(r[colmap["high"]], conv) if "high" in colmap else None
                lo = _num(r[colmap["low"]], conv) if "low" in colmap else None
                o, h, lo = (x * f if x is not None else None for x in (o, h, lo))
                v = _vol(r[colmap["volume"]], conv) if "volume" in colmap and colmap["volume"] < len(r) else 0
            except (IndexError, TypeError, ValueError, ZeroDivisionError):
                continue
            if not tk or not d or c is None:
                continue
            o = o or c
            h = max(x for x in (h, o, c) if x is not None)
            lo = min(x for x in (lo, o, c) if x is not None)
            parsed.append([tk, d, o, h, lo, c, v or 0])
        # đơn vị nghìn đồng → đồng: quyết định cho CẢ mã (trung vị), không theo từng dòng
        by_ticker: dict[str, list[list]] = {}
        for row in parsed:
            by_ticker.setdefault(row[0], []).append(row)
        for tk, rows in by_ticker.items():
            closes = sorted(r[5] for r in rows)
            if tk not in {"VNINDEX", "VN30", "HNXINDEX", "UPCOMINDEX", "HNX30", "VN100"} and closes[len(closes) // 2] < 500:
                for r in rows:
                    r[2:6] = [x * 1_000 for x in r[2:6]]
            rows_out.extend(tuple(r) for r in rows)
    return rows_out


def import_price_file(data: bytes, filename: str, default_ticker: str | None = None) -> dict:
    rows = parse_price_file(data, filename, default_ticker)
    if not rows:
        raise ValueError("Không đọc được dữ liệu giá. Cần các cột: Mã, Ngày, Mở cửa, Cao nhất, Thấp nhất, "
                         "Đóng cửa, Khối lượng (hoặc định dạng Amibroker <Ticker>,<DTYYYYMMDD>,...).")
    db = get_db()
    with db._lock:  # noqa: SLF001
        db._conn.executemany(
            "INSERT OR REPLACE INTO imported_prices(ticker,date,open,high,low,close,volume) VALUES(?,?,?,?,?,?,?)",
            rows)
        db._conn.commit()
    db.execute("DELETE FROM cache WHERE key LIKE 'hist:%'")
    counts: dict[str, int] = {}
    for r in rows:
        counts[r[0]] = counts.get(r[0], 0) + 1
    dates = sorted(r[1] for r in rows)
    return {"rows": len(rows), "tickers": counts, "from": dates[0], "to": dates[-1]}


def imported_summary() -> list[dict]:
    return get_db().query("SELECT ticker, COUNT(*) AS rows, MIN(date) AS first, MAX(date) AS last "
                          "FROM imported_prices GROUP BY ticker ORDER BY ticker")


def delete_imported(ticker: str) -> None:
    get_db().execute("DELETE FROM imported_prices WHERE ticker=?", (ticker.upper(),))
    get_db().execute("DELETE FROM cache WHERE key LIKE 'hist:%'")
