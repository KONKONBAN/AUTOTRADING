import pytest
from conftest import login
from fastapi.testclient import TestClient
from make_sample_xlsx import vietstock_like

from app.data import excel_import
from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield login(c)


def test_excel_parser_vietstock_layout():
    r = excel_import.parse_financial_file(vietstock_like(), "fpt.xlsx")
    y = {x["year"]: x for x in r["years"]}[2024]
    assert r["unit_detected"] == "trieu dong"
    assert y["revenue"] == pytest.approx(62_800)          # 62.800.000 triệu = 62.800 tỷ
    assert y["net_income"] == pytest.approx(7_800)        # ưu tiên LNST cổ đông công ty mẹ
    assert y["total_debt"] == pytest.approx(14_000)       # vay ngắn + dài hạn
    assert y["capex"] == pytest.approx(4_000)             # luôn dương
    assert y["shares"] == pytest.approx(1_470)            # vốn góp 14.700 tỷ / 10.000đ


def test_template_roundtrip():
    import io

    from openpyxl import load_workbook
    data = excel_import.financial_template("FPT")
    wb = load_workbook(io.BytesIO(data))
    ws = wb["BCTC"]
    # file mẫu KHÔNG được điền sẵn số liệu DEMO
    assert all(ws.cell(r, c).value is None for r in range(5, 24) for c in range(4, 9))
    values = {"revenue": 1000, "net_income": 100, "total_assets": 2000, "equity": 800, "shares": 50}
    for r in range(5, 24):
        key = ws.cell(r, 2).value
        if key in values:
            for c in range(4, 9):
                ws.cell(r, c, values[key])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = excel_import.parse_financial_file(buf.getvalue(), "t.xlsx")
    assert len(parsed["years"]) == 5
    assert parsed["years"][-1]["revenue"] == 1000 and parsed["years"][-1]["shares"] == 50


def test_price_csv_formats():
    ami = b"<Ticker>,<DTYYYYMMDD>,<Open>,<High>,<Low>,<Close>,<Volume>\nAAA,20250102,10,11,9,10.5,100\n"
    rows = excel_import.parse_price_file(ami, "a.csv")
    assert rows[0][:2] == ("AAA", "2025-01-02") and rows[0][5] == 10_500
    cafef = "Ngày,Giá điều chỉnh,Giá đóng cửa,Khối lượng khớp lệnh,Giá mở cửa,Giá cao nhất,Giá thấp nhất\n03/01/2025,21.5,21.5,1200,20.5,22,20\n"
    rows = excel_import.parse_price_file(cafef.encode(), "b.csv", "BBB")
    assert rows[0][0] == "BBB" and rows[0][5] == 21_500


def test_core_endpoints(client):
    for url in ["/api/health", "/api/status", "/api/market", "/api/universe", "/api/stocks/FPT/analysis",
                "/api/stocks/HPG/technical?start=2025-06-01&end=2025-12-31", "/api/financials/FPT",
                "/api/portfolio", "/api/autotrader/config", "/api/knowledge"]:
        assert client.get(url).status_code == 200, url
    a = client.get("/api/stocks/FPT/analysis").json()
    assert 0 <= a["composite"]["score"] <= 100
    assert a["data_sources"]["prices"] == "demo"


def test_paper_trading_rules(client):
    client.post("/api/portfolio/reset", json={"initial_cash": 100_000_000})
    ok = client.post("/api/orders", json={"ticker": "VCB", "side": "BUY", "quantity": 100})
    assert ok.status_code == 200 and ok.json()["fee"] > 0
    odd = client.post("/api/orders", json={"ticker": "VCB", "side": "BUY", "quantity": 150})
    assert odd.status_code == 400 and "lô chẵn" in odd.json()["detail"]
    t2 = client.post("/api/orders", json={"ticker": "VCB", "side": "SELL", "quantity": 100})
    assert t2.status_code == 400 and "T+2" in t2.json()["detail"]
    broke = client.post("/api/orders", json={"ticker": "FPT", "side": "BUY", "quantity": 100_000})
    assert broke.status_code == 400


def test_manual_financials_override(client):
    rows = [{"year": 2025, "revenue": 70000, "net_income": 9000, "total_assets": 85000, "equity": 40000, "shares": 1700}]
    assert client.put("/api/financials/FPT/manual", json={"rows": rows}).status_code == 200
    fin = client.get("/api/financials/FPT").json()
    assert 2025 in fin["manual_years"] and fin["source"].startswith("manual")
    bad = client.put("/api/financials/FPT/manual", json={"rows": [{"year": 2024, "revenue": 1}]})
    assert bad.status_code == 400


def test_autotrader_suggest_and_approve(client):
    client.post("/api/portfolio/reset", json={"initial_cash": 500_000_000})
    client.put("/api/autotrader/config", json={"mode": "suggest", "min_score": 50, "min_technical": 40})
    run = client.post("/api/autotrader/run", json={}).json()
    assert run["mode"] == "suggest"
    pending = client.get("/api/proposals").json()
    if pending:
        res = client.post(f"/api/proposals/{pending[0]['id']}/approve").json()
        assert res["status"] in {"executed", "failed"}


def test_backtest_and_optimizer(client):
    bt = client.post("/api/backtest", json={"tickers": ["FPT", "VCB", "HPG", "MWG"], "min_score": 60}).json()
    assert "metrics" in bt and len(bt["equity_curve"]) > 100
    opt = client.post("/api/portfolio/optimize", json={"tickers": ["FPT", "VCB", "HPG", "MWG", "GAS"], "max_weight": 0.3}).json()
    assert sum(w["weight"] for w in opt["weights"]) == pytest.approx(1, abs=1e-6)
    assert max(w["weight"] for w in opt["weights"]) <= 0.3 + 1e-6


def test_agent_offline_and_ai_endpoints_without_key(client):
    r = client.post("/api/agent/chat", data={"message": "Phân tích VCB"}).json()
    assert r["offline"] and "VCB" in r["reply"]
    chart = client.post("/api/charts/analyze", files={"files": ("a.png", b"\x89PNG")})
    assert chart.status_code == 400 and "ANTHROPIC_API_KEY" in chart.json()["detail"]


def test_finance_endpoints(client):
    plan = client.post("/api/finance/plan", json={"monthly_income": 30e6, "monthly_expenses": 18e6,
                                                  "goals": [{"name": "Nhà", "amount": 2e9, "years": 8}]}).json()
    assert plan["summary"]["surplus"] == 12e6 and plan["goals"][0]["monthly_saving"] > 0
    calc = client.post("/api/finance/calc", json={"calculator": "bond_ytm", "params": {
        "price": 1_035_000, "face": 1_000_000, "coupon_rate": 0.1, "years": 2}}).json()
    assert calc["result"] == pytest.approx(0.0804, abs=1e-4)
