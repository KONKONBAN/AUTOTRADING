"""Kiểm thử hồi quy cho các lỗi đã sửa sau đợt review trước khi tạo PR."""
import pytest
from fastapi.testclient import TestClient

from app.agent import agent as agent_mod
from app.analysis import fundamental
from app.data import excel_import, fundamentals, providers
from app.finance import calculators, planner
from app.main import app
from app.services import sector_multiples
from app.trading import broker as broker_mod
from app.trading import engine


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


# ------------------------------------------------------------------ dữ liệu / Excel
def test_number_conventions():
    assert excel_import._num("1.234", "vn") == 1234
    assert excel_import._num("1,5", "vn") == 1.5
    assert excel_import._num("1.234.567,8") == pytest.approx(1_234_567.8)
    assert excel_import._num("1,234,567.8") == pytest.approx(1_234_567.8)
    assert excel_import._num("(1.200)", "vn") == -1200
    assert excel_import.detect_convention([["1.234.567", "2.345,6"]]) == "vn"
    assert excel_import.detect_convention([["1,234,567", "2,345.6"]]) == "en"


def test_unit_detection_numeric_multipliers():
    assert excel_import._detect_unit([["Đơn vị: 1.000.000 VND"]])[0] == pytest.approx(1e-3)
    assert excel_import._detect_unit([["Đơn vị tính: 1.000 đồng"]])[0] == pytest.approx(1e-6)
    assert excel_import._detect_unit([["Đơn vị: tỷ đồng"]])[0] == pytest.approx(1.0)


def test_numeric_data_rows_are_not_years():
    assert excel_import._row_years(["Doanh thu", 1900, 2000, 2100]) == {}
    assert excel_import._row_years(["Chỉ tiêu", 2022, 2023, 2024]) == {1: 2022, 2: 2023, 3: 2024}


def test_quarterly_only_file_is_rejected_with_clear_message():
    csv = ("Chỉ tiêu,Q1/2024,Q2/2024,Q3/2024,Q4/2024\n"
           "Doanh thu thuần,100,110,120,130\nLợi nhuận sau thuế,10,11,12,13\n").encode()
    with pytest.raises(ValueError, match="QUÝ"):
        excel_import.parse_financial_file(csv, "q.csv")


def test_manual_rows_never_mixed_with_demo():
    t = "VNM"
    fundamentals.delete_manual(t)
    fundamentals.save_manual(t, [{"year": 2024, "revenue": 60_000, "net_income": 9_000,
                                  "total_assets": 50_000, "equity": 35_000, "total_debt": 0}])
    fin = fundamentals.get_financials(t)
    assert fin["source"] == "manual" and [r["year"] for r in fin["rows"]] == [2024]
    row = fin["rows"][0]
    assert row["total_debt"] == 0 and "total_debt" not in row["estimated_fields"]  # 0 hợp lệ, không ước tính
    assert row["shares"] is None and row["eps"] is None  # không lấy số cổ phiếu từ DEMO
    fundamentals.delete_manual(t)


def test_demo_prices_stable_across_lengths():
    a = providers.get_history("FPT", days=200)
    b = providers.get_history("FPT", days=600)
    common = a.index.intersection(b.index)
    assert len(common) > 100
    assert (a.loc[common, "close"] == b.loc[common, "close"]).all()


# ------------------------------------------------------------------ phân tích / tài chính
def test_negative_equity_has_no_roe():
    rows = [{"year": 2024, "revenue": 100, "net_income": -20, "total_assets": 80, "equity": -10,
             "total_liabilities": 90, "interest_expense": 5}]
    r = fundamental.ratios_by_year(rows)[0]
    assert r["negative_equity"] and r["roe"] is None and r["lev"] is None


def test_irr_root_on_grid_point_and_textbook():
    assert calculators.irr([-100, 100]) == 0.0
    assert calculators.irr([-1, -0.1, -0.5, 0.8, 1.0]) == pytest.approx(0.0417, abs=5e-4)


def test_lease_salvage_cashflow():
    res = calculators.lease_vs_buy(1000, 250, 4, 0.10, tax_rate=0.2, salvage=200)
    dep = (1000 - 200) / 4
    assert res["cashflows"][-1] == pytest.approx(-dep * 0.2 - 200)


def test_planner_builds_emergency_fund_first():
    p = planner.plan(30_000_000, 20_000_000, current_savings=0, risk_profile="can_bang")
    eb = p["emergency_build"]
    assert eb["months"] == 12 and eb["monthly"] == pytest.approx(10_000_000)
    assert eb["invest_during_build"] == pytest.approx(0)
    year1 = p["projection"][0]
    assert year1["years"] == 1 and year1["invested"] == pytest.approx(0)
    assert year1["emergency_fund"] >= 120_000_000


def test_sector_multiples_exclude_self():
    from app.data import universe
    sector = universe.info("VCB")["sector"]
    with_self = sector_multiples(sector)
    without = sector_multiples(sector, exclude="VCB")
    assert without["n"] <= with_self["n"]


# ------------------------------------------------------------------ giao dịch
def test_limit_orders_and_typos():
    b = broker_mod.get_broker()
    b.reset(500_000_000)
    last = float(providers.get_history("FPT", days=60)["close"].iloc[-1])
    with pytest.raises(broker_mod.OrderError, match="thấp hơn"):
        b.place_order("FPT", "BUY", 100, last * 0.97)
    with pytest.raises(broker_mod.OrderError, match="danh sách"):
        b.place_order("FTP", "BUY", 100)
    res = b.place_order(" fpt ", "BUY", 100, last * 1.03)
    assert res["price"] == last  # luôn khớp tại giá thị trường
    b.reset(500_000_000)


def test_proposal_cannot_be_approved_twice():
    broker_mod.get_broker().reset(500_000_000)
    pid = engine.create_proposal("hpg", "buy", 100, None, "test")["proposal_id"]
    first = engine.decide_proposal(pid, True)
    assert first["status"] == "executed"
    with pytest.raises(ValueError):
        engine.decide_proposal(pid, True)
    with pytest.raises(ValueError):
        engine.create_proposal("XXXX", "BUY", 100, None, "typo")
    broker_mod.get_broker().reset(500_000_000)


# ------------------------------------------------------------------ API
def test_technical_endpoint_rejects_bad_ranges(client):
    assert client.get("/api/stocks/FPT/technical?start=abc").status_code == 400
    assert client.get("/api/stocks/FPT/technical?start=2025-06-01&end=2025-01-01").status_code == 400


def test_unknown_api_path_is_404_not_spa(client):
    assert client.get("/api/khong-ton-tai").status_code == 404


def test_upload_rejects_unsupported_file(client):
    r = client.post("/api/financials/upload", files={"file": ("bctc.pdf", b"%PDF-1.4 ...", "application/pdf")})
    assert r.status_code == 400


# ------------------------------------------------------------------ AI agent
def test_dangling_tool_use_is_closed():
    msgs = [{"role": "user", "content": "hi"},
            {"role": "assistant", "content": [{"type": "text", "text": "..."},
                                              {"type": "tool_use", "id": "t1", "name": "x", "input": {}}]}]
    agent_mod._close_dangling_tool_uses(msgs)
    assert msgs[-1]["role"] == "user" and msgs[-1]["content"][0]["tool_use_id"] == "t1"
    agent_mod._append_user(msgs, [{"type": "text", "text": "tiếp"}])
    assert [m["role"] for m in msgs] == ["user", "assistant", "user"]


def test_system_prompt_is_stable():
    assert agent_mod._system() == agent_mod._system()
    assert "Ngày hiện tại" not in agent_mod._system()[0]["text"]


# ------------------------------------------------------------------ vòng kiểm chứng thứ hai
def test_decimal_numbers_not_read_as_thousands():
    assert excel_import._num("44010.123") == pytest.approx(44010.123)
    assert excel_import._num("0.125") == pytest.approx(0.125)
    assert excel_import._num("0,125") == pytest.approx(0.125)
    assert excel_import._num("12.345") == 12345          # không rõ quy ước: ưu tiên kiểu VN
    assert excel_import._num("1.234", "en") == pytest.approx(1.234)
    assert excel_import._num("1,234", "vn") == pytest.approx(1.234)


@pytest.mark.parametrize("label", ["2024Q1", "4Q2024", "Q1/2024", "Quý 4/2024", "2024H1", "6T2024", "9 tháng 2024", "2025f"])
def test_quarter_and_forecast_labels_are_not_years(label):
    assert excel_import._year_of(label) is None


def test_header_with_cagr_summary_cell_still_parses():
    assert excel_import._row_years(["Chỉ tiêu", 2020, 2021, 2022, 2023, 2024, "CAGR 2020-2024"]) == \
        {1: 2020, 2: 2021, 3: 2022, 4: 2023, 5: 2024}
    csv = ("Đơn vị: tỷ đồng\nChỉ tiêu,2022,2023,2024,CAGR 2022-2024\nDoanh thu thuần,100,110,121,10%\n"
           "Lợi nhuận sau thuế,10,11,12,9%\nTổng tài sản,200,210,220,5%\nVốn chủ sở hữu,120,125,130,4%\n").encode()
    assert [y["year"] for y in excel_import.parse_financial_file(csv, "c.csv")["years"]] == [2022, 2023, 2024]


def test_corrupt_xlsx_gives_format_error(client):
    import io
    import zipfile

    from make_sample_xlsx import vietstock_like
    src = zipfile.ZipFile(io.BytesIO(vietstock_like()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            z.writestr(item, data[: len(data) // 2] if item.filename.startswith("xl/worksheets/") else data)
    with pytest.raises(excel_import.FileFormatError):
        excel_import.parse_financial_file(buf.getvalue(), "broken.xlsx")
    r = client.post("/api/financials/upload", files={"file": ("broken.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 400


def test_template_keeps_estimated_fields_blank():
    import io

    from openpyxl import load_workbook
    t = "VNM"
    fundamentals.delete_manual(t)
    fundamentals.save_manual(t, [{"year": 2024, "revenue": 60_000, "net_income": 9_000,
                                  "total_assets": 50_000, "equity": 35_000, "shares": 2_090}])
    ws = load_workbook(io.BytesIO(excel_import.financial_template(t)))["BCTC"]
    vals = {ws.cell(r, 2).value: ws.cell(r, 4).value for r in range(5, 24)}
    assert vals["revenue"] == 60_000 and vals["shares"] == 2_090
    assert vals["total_debt"] is None and vals["pretax_income"] is None and vals["operating_cf"] is None
    assert vals["interest_expense"] is None  # bỏ trống → không ghi lại thành 0 "thật"
    fundamentals.delete_manual(t)


def test_shares_from_owner_capital_flag():
    csv = ("Chỉ tiêu,2023,2024\nDoanh thu thuần,50000,60000\nLợi nhuận sau thuế,5000,6000\n"
           "Tổng tài sản,90000,100000\nVốn chủ sở hữu,40000,45000\nVốn góp của chủ sở hữu,10000,10000\n").encode()
    assert excel_import.parse_financial_file(csv, "s.csv")["shares_from_capital"] is True


def test_pe_gordon_uses_steady_state_payout():
    from app.analysis import valuation
    rows = [{"year": 2024, "eps": 4000, "bvps": 22000, "dps": 700, "shares": 100, "net_income": 400}]
    ratios = [{"roe": 0.184}]
    v = valuation.valuation(rows, ratios, price=41_850, beta=1.0, sector="Ngân hàng")
    pe = next(m for m in v["models"] if m["model"].startswith("P/E hợp lý"))
    assert pe["value"] / 4000 > 6  # không còn bị kẹp ở sàn 5x do dùng tỷ lệ chi trả hiện tại thấp


def test_negative_equity_scores_zero_roe_stability():
    rows = [{"year": y, "revenue": 1000, "net_income": ni, "total_assets": 2000, "equity": e,
             "total_liabilities": 2000 - e} for y, ni, e in
            [(2021, 120, 1000), (2022, 125, 1100), (2023, 130, 1200), (2024, -1400, -200)]]
    q = fundamental.score_quality(fundamental.ratios_by_year(rows), {"score": None}, "Công nghiệp")
    assert next(p for p in q["parts"] if p["name"] == "Ổn định ROE")["score"] == 0


def test_single_year_has_no_cagr():
    rows = [{"year": 2024, "revenue": 1000, "net_income": 100, "total_assets": 2000, "equity": 900,
             "total_liabilities": 1100}]
    q = fundamental.score_quality(fundamental.ratios_by_year(rows), {"score": None}, "Công nghiệp")
    assert q["revenue_cagr"] is None and q["ni_cagr"] is None
    assert not any("CAGR" in p["name"] for p in q["parts"])


def _settled_lot(ticker: str, qty: int, price: float):
    from app.db import get_db
    get_db().execute("INSERT INTO lots(ticker, quantity, price, buy_date) VALUES(?,?,?,?)",
                     (ticker, qty, price, "2020-01-02"))


def test_priced_sell_proposal_not_sold_below_approved_price():
    broker_mod.get_broker().reset(500_000_000)
    last = broker_mod.quote("FPT")["last"]
    _settled_lot("FPT", 1000, last * 0.9)
    pid = engine.create_proposal("FPT", "SELL", 1000, last * 1.03, "Chốt lời", source="agent")["proposal_id"]
    assert engine.decide_proposal(pid, True)["status"] == "failed"   # thị trường thấp hơn giá đã duyệt
    pid = engine.create_proposal("FPT", "SELL", 1000, last * 1.03, "Cắt lỗ", source="autotrader")["proposal_id"]
    assert engine.decide_proposal(pid, True)["status"] == "executed"  # thoát vị thế của Auto-Trader: lệnh thị trường
    broker_mod.get_broker().reset(500_000_000)


def test_buy_proposal_near_ceiling_is_clamped():
    broker_mod.get_broker().reset(500_000_000)
    q = broker_mod.quote("HPG")
    pid = engine.create_proposal("HPG", "BUY", 100, q["limits"].ceiling, "test", source="agent")["proposal_id"]
    res = engine.decide_proposal(pid, True)
    assert res["status"] == "executed", res
    broker_mod.get_broker().reset(500_000_000)


def test_buy_at_market_price_with_float_noise(monkeypatch):
    b = broker_mod.get_broker()
    b.reset(500_000_000)
    real = broker_mod.quote

    def noisy(t):
        q = real(t)
        return {**q, "last": round(q["last"] + 1e-7, 2) + 1e-9}
    monkeypatch.setattr(broker_mod, "quote", noisy)
    last = real("FPT")["last"]
    assert b.place_order("FPT", "BUY", 100, last)["status"] == "FILLED"
    b.reset(500_000_000)


def test_timezone_aware_dates_accepted(client):
    r = client.get("/api/stocks/FPT/technical", params={"start": "2025-01-01T00:00:00+07:00",
                                                        "end": "2025-12-31T00:00:00+07:00"})
    assert r.status_code == 200


def test_agent_chat_serialized_per_session(monkeypatch):
    import threading
    import time as _t
    active, peak = [0], [0]

    def fake(message, session_id, attachments):
        active[0] += 1
        peak[0] = max(peak[0], active[0])
        _t.sleep(0.05)
        active[0] -= 1
        return {"session_id": session_id}
    monkeypatch.setattr(agent_mod, "_chat", fake)
    ts = [threading.Thread(target=agent_mod.chat, args=(f"q{i}", "same-session")) for i in range(4)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert peak[0] == 1


# ------------------------------------------------------------------ vòng kiểm chứng thứ ba
def test_pe_gordon_monotonic_in_roe():
    from app.analysis import valuation

    def pe(roe):
        rows = [{"year": 2024, "eps": 1000, "bvps": 10_000, "dps": 0, "shares": 100, "net_income": 100}]
        v = valuation.valuation(rows, [{"roe": roe}], price=10_000, beta=1.0, sector="Ngân hàng")
        m = next((m for m in v["models"] if m["model"].startswith("P/E hợp lý")), None)
        return m["value"] / 1000 if m else None
    assert pe(None) is None and pe(-0.1) is None   # vốn chủ âm / ROE âm: không dùng mô hình
    seq = [pe(x) for x in (0.02, 0.035, 0.04, 0.05, 0.08, 0.10, 0.15, 0.25)]
    assert all(b >= a - 1e-9 for a, b in zip(seq, seq[1:])), seq


def test_english_three_decimal_file():
    csv = ("Unit: billion VND\nItem,2022,2023,2024\nNet revenue,44010.123,52618.456,62849.789\n"
           "Net profit after tax,950.500,1020.250,1100.125\nTotal assets,61697.123,65000.456,70000.789\n"
           "Total equity,25000.000,27000.500,29000.250\n").encode()
    r = excel_import.parse_financial_file(csv, "en.csv")
    assert r["number_format"] == "en"
    assert [round(y["net_income"], 3) for y in r["years"]] == [950.5, 1020.25, 1100.125]


@pytest.mark.parametrize("label", ["1H2025", "2H2024", "2025 1H", "Bán niên 2024"])
def test_half_year_labels_are_not_years(label):
    assert excel_import._year_of(label) is None


def test_year_header_keeps_all_runs():
    assert excel_import._row_years(["Chỉ tiêu", 2014, 2015, 2016, 2023, 2024]) == \
        {1: 2014, 2: 2015, 3: 2016, 4: 2023, 5: 2024}


@pytest.mark.parametrize("member", ["xl/styles.xml", "xl/workbook.xml"])
def test_corrupt_workbook_parts_give_format_error(member):
    import io
    import zipfile

    from make_sample_xlsx import vietstock_like
    src = zipfile.ZipFile(io.BytesIO(vietstock_like()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            z.writestr(item, data[: len(data) // 2] if item.filename == member else data)
    with pytest.raises(excel_import.FileFormatError):
        excel_import.parse_financial_file(buf.getvalue(), "broken.xlsx")


def test_blank_current_liabilities_is_estimated_not_zero():
    row = fundamentals._complete_row({"year": 2024, "revenue": 100, "net_income": 10, "total_assets": 200,
                                      "equity": 100, "current_assets": 120})
    assert row["current_liabilities"] == pytest.approx(60) and "current_liabilities" in row["estimated_fields"]
    row = fundamentals._complete_row({"year": 2024, "revenue": 100, "net_income": 10, "total_assets": 200,
                                      "equity": 100})
    assert row["current_liabilities"] == 0 and "current_liabilities" in row["defaulted_fields"]
